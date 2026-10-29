"""Phase 0 – Inventur und Rohextraktion (rein mechanisch, kein LLM).

Für jede Datei in input/raw/:
  * SHA-256 → identische Dateien werden als Duplikat erkannt
  * jedes Blatt wird zellgenau nach work/extrakt/<blatt_id>.jsonl geschrieben,
    inklusive Füllfarbe (Theme-Farben aufgelöst), Schriftfarbe, Durchstreichung,
    Kommentar, verbundene/ausgeblendete Zellen
  * kompakte Vorschau work/extrakt/<blatt_id>.vorschau.md für den Agenten
  * Layout-Signatur (Kopfzeile) → Blätter gleicher Vorlage bekommen dasselbe Profil
  * Inhalts-Hash → identische Blätter in verschiedenen Dateien werden erkannt

Manuell gepflegte Felder im Manifest (fachbereich, nav_tabelle, profil, notiz)
bleiben bei erneutem Lauf erhalten. Danach wird automatisch `queue.py sync` ausgeführt.

Aufruf:  python scripts/p0_extract.py [--force]
"""
from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

from common import (EXTRAKT, MANIFEST, RAW, append_log, load_yaml, norm, now,
                    resolve_color, save_yaml, sha256_file, theme_colors, write_text)

LESBAR = {".xlsx", ".xlsm"}
EDITIERBAR_DATEI = ("fachbereich", "notiz")
EDITIERBAR_BLATT = ("profil", "nav_tabelle", "nav_tabelle_quelle", "fachbereich", "notiz")
VORSCHAU_ZEILEN = 40


def _wert(v):
    if isinstance(v, bool):
        return ("WAHR" if v else "FALSCH"), "bool"
    if isinstance(v, (int, float)):
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return str(v), "zahl"
    if hasattr(v, "isoformat"):
        return v.isoformat(), "datum"
    return str(v), "text"


def extract_sheet(ws, ws_val, theme) -> tuple[list[dict], dict]:
    merged = {}
    for rng in ws.merged_cells.ranges:
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                merged[(row, col)] = str(rng)
    hidden_cols = set()
    for d in ws.column_dimensions.values():
        if d.hidden and d.min and d.max:
            hidden_cols.update(range(d.min, d.max + 1))
    hidden_rows = {r for r, d in ws.row_dimensions.items() if d.hidden}

    cells = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c, MergedCell):
                continue
            raw = c.value
            kommentar = c.comment.text.strip() if c.comment else None
            leer = raw is None or (isinstance(raw, str) and raw.strip() == "")
            if leer and not kommentar:
                continue
            rec = {"zelle": c.coordinate, "zeile": c.row, "spalte": get_column_letter(c.column)}
            if not leer:
                if isinstance(raw, str) and raw.startswith("="):
                    rec["formel"] = raw
                    wert, typ = _wert(ws_val.cell(c.row, c.column).value or "")
                    rec["wert"], rec["typ"] = wert, "formel"
                else:
                    rec["wert"], rec["typ"] = _wert(raw)
            else:
                rec["wert"], rec["typ"] = "", "leer"
            if c.fill is not None and c.fill.fill_type:
                f = resolve_color(c.fill.fgColor, theme)
                if f:
                    rec["fuellung"] = f
            if c.font is not None:
                fc = resolve_color(c.font.color, theme)
                if fc and fc not in ("#000000",):
                    rec["schriftfarbe"] = fc
                if c.font.strike:
                    rec["durchgestrichen"] = True
                if c.font.bold:
                    rec["fett"] = True
            if kommentar:
                rec["kommentar"] = kommentar
            if (c.row, c.column) in merged:
                rec["verbunden"] = merged[(c.row, c.column)]
            if c.column in hidden_cols or c.row in hidden_rows:
                rec["versteckt"] = True
            cells.append(rec)

    meta = {
        "zellen": len([c for c in cells if c["typ"] != "leer"]),
        "zeilen": max((c["zeile"] for c in cells), default=0),
        "sichtbarkeit": ws.sheet_state,
        "versteckte_spalten": sorted(get_column_letter(i) for i in hidden_cols) or None,
        "versteckte_zeilen": len(hidden_rows) or None,
        "verbundene_bereiche": len(ws.merged_cells.ranges) or None,
    }
    return cells, meta


def signatur(cells: list[dict]) -> str:
    """Kopfzeile = erste Zeile (unter den ersten 20) mit >= 2 Textzellen."""
    by_row: dict[int, list[dict]] = {}
    for c in cells:
        if c["typ"] != "leer":
            by_row.setdefault(c["zeile"], []).append(c)
    for r in sorted(by_row)[:20]:
        texte = [c for c in by_row[r] if c["typ"] == "text"]
        if len(texte) >= 2:
            tokens = [f'{c["spalte"]}:{norm(c["wert"])}' for c in sorted(by_row[r], key=lambda x: (len(x["spalte"]), x["spalte"]))]
            return "L" + hashlib.sha1("|".join(tokens).encode()).hexdigest()[:8]
    return "L" + hashlib.sha1(b"ohne_kopf").hexdigest()[:8]


def inhalt_hash(cells: list[dict]) -> str:
    pairs = sorted((c["zelle"], c["wert"]) for c in cells if c["typ"] != "leer")
    return hashlib.sha1(json.dumps(pairs, ensure_ascii=False).encode()).hexdigest()[:16]


def vorschau(blatt: dict, datei: str, cells: list[dict]) -> str:
    """Kompakte, token-sparsame Sicht für den Agenten."""
    farben = Counter(c.get("fuellung") for c in cells if c.get("fuellung"))
    fkurz = {f: f"f{i + 1}" for i, (f, _) in enumerate(farben.most_common())}
    sfarben = Counter(c.get("schriftfarbe") for c in cells if c.get("schriftfarbe"))
    skurz = {f: f"s{i + 1}" for i, (f, _) in enumerate(sfarben.most_common())}

    spalten = sorted({c["spalte"] for c in cells}, key=lambda s: (len(s), s))
    zeilen = sorted({c["zeile"] for c in cells})
    grid = {(c["zeile"], c["spalte"]): c for c in cells}

    out = [f"# Vorschau {blatt['blatt_id']}", "",
           f"- Datei: `{datei}`", f"- Blatt: `{blatt['name']}` ({blatt['meta']['sichtbarkeit']})",
           f"- Signatur: `{blatt['signatur']}`  Zellen: {blatt['meta']['zellen']}  Zeilen: {blatt['meta']['zeilen']}"]
    for k in ("versteckte_spalten", "versteckte_zeilen", "verbundene_bereiche"):
        if blatt["meta"].get(k):
            out.append(f"- {k}: {blatt['meta'][k]}")
    out += ["", "## Legende Formatierung",
            "`·f1` = Füllfarbe, `·s1` = Schriftfarbe, `~~x~~` = durchgestrichen, `†` = Kommentar, `(v)` = versteckt", ""]
    for f, n in farben.most_common():
        beispiele = [c["zelle"] for c in cells if c.get("fuellung") == f][:6]
        out.append(f"- {fkurz[f]} = Füllung {f}: {n} Zellen (z. B. {', '.join(beispiele)})")
    for f, n in sfarben.most_common():
        out.append(f"- {skurz[f]} = Schrift {f}: {n} Zellen")
    kommentare = [c for c in cells if c.get("kommentar")]
    if kommentare:
        out += ["", "## Kommentare"] + [f"- {c['zelle']}: {c['kommentar']}" for c in kommentare[:30]]

    out += ["", "## Spaltenprofil", "", "| Spalte | belegt | häufigste Werte |", "|---|---|---|"]
    for s in spalten:
        werte = Counter(c["wert"] for c in cells if c["spalte"] == s and c["typ"] != "leer")
        top = ", ".join(f"`{w[:25]}`×{n}" for w, n in werte.most_common(6))
        out.append(f"| {s} | {sum(werte.values())} | {top} |")

    def zelle(r, s):
        c = grid.get((r, s))
        if not c:
            return ""
        t = c["wert"].replace("|", "/").replace("\n", " ")[:30]
        if c.get("durchgestrichen"):
            t = f"~~{t}~~"
        if c.get("fuellung"):
            t += f" ·{fkurz[c['fuellung']]}"
        if c.get("schriftfarbe"):
            t += f" ·{skurz[c['schriftfarbe']]}"
        if c.get("kommentar"):
            t += " †"
        if c.get("versteckt"):
            t += " (v)"
        return t

    zeige = zeilen[:VORSCHAU_ZEILEN]
    rest = zeilen[VORSCHAU_ZEILEN:]
    out += ["", f"## Raster (erste {len(zeige)} von {len(zeilen)} belegten Zeilen)", "",
            "| Z | " + " | ".join(spalten) + " |", "|---" * (len(spalten) + 1) + "|"]
    for r in zeige:
        out.append(f"| {r} | " + " | ".join(zelle(r, s) for s in spalten) + " |")
    if rest:
        out.append(f"| … | {len(rest) - min(5, len(rest))} Zeilen ausgelassen |")
        for r in rest[-5:]:
            out.append(f"| {r} | " + " | ".join(zelle(r, s) for s in spalten) + " |")
    out.append("")
    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true", help="alle Dateien neu extrahieren")
    args = ap.parse_args()

    manifest = load_yaml(MANIFEST, {"dateien": {}})
    alt = manifest.get("dateien", {})
    neu: dict = {}
    EXTRAKT.mkdir(parents=True, exist_ok=True)

    dateien = sorted(p for p in RAW.rglob("*") if p.is_file()
                     and not p.name.startswith(("~$", ".")) and p.name != "LIESMICH.md")
    inhalt_index: dict[str, str] = {}
    datei_index: dict[str, str] = {}
    zaehler = Counter()

    for p in dateien:
        rel = p.relative_to(RAW).as_posix()
        sha = sha256_file(p)
        # ID hängt am Pfad, nicht am Inhalt: überarbeitet ein Fachbereich seine Datei,
        # bleiben IDs und manuelle Felder erhalten, und die Queue öffnet die Einheit neu.
        fid = hashlib.sha1(rel.encode()).hexdigest()[:10]
        vorher = alt.get(fid, {})
        eintrag = {"pfade": [rel], "sha256": sha, "format": p.suffix.lower()}
        if sha in datei_index:  # byte-identische Datei unter anderem Namen
            eintrag["duplikat_von"] = datei_index[sha]
            eintrag["notiz"] = vorher.get("notiz")
            neu[fid] = eintrag
            zaehler["duplikat_datei"] += 1
            continue
        datei_index[sha] = fid
        teile = rel.split("/")
        eintrag["fachbereich"] = vorher.get("fachbereich", teile[0] if len(teile) > 1 else None)
        eintrag["notiz"] = vorher.get("notiz")

        if p.suffix.lower() not in LESBAR:
            eintrag["lesbar"] = False
            eintrag["hinweis"] = "Format nicht unterstützt – bitte als .xlsx speichern (z. B. soffice --headless --convert-to xlsx)"
            neu[fid] = eintrag
            zaehler["nicht_lesbar"] += 1
            continue

        alte_blaetter = {b["blatt_id"]: b for b in vorher.get("blaetter", [])}
        unveraendert = (not args.force and vorher.get("sha256") == sha and alte_blaetter
                        and all((EXTRAKT / f"{bid}.jsonl").exists() for bid in alte_blaetter))
        if unveraendert:
            eintrag["blaetter"] = list(alte_blaetter.values())
            for b in eintrag["blaetter"]:
                b.pop("duplikat_von", None)  # wird unten neu bestimmt
        else:
            wb = openpyxl.load_workbook(p, data_only=False)
            wbv = openpyxl.load_workbook(p, data_only=True)
            theme = theme_colors(wb)
            eintrag["blaetter"] = []
            for i, ws in enumerate(wb.worksheets):
                bid = f"{fid}-{i + 1}"
                cells, meta = extract_sheet(ws, wbv.worksheets[i], theme)
                b = {"blatt_id": bid, "name": ws.title, "meta": meta,
                     "signatur": signatur(cells), "inhalt_hash": inhalt_hash(cells)}
                ab = alte_blaetter.get(bid, {})
                for k in EDITIERBAR_BLATT:
                    b[k] = ab.get(k)
                b["profil"] = b["profil"] or b["signatur"]
                with open(EXTRAKT / f"{bid}.jsonl", "w", encoding="utf-8") as f:
                    for c in cells:
                        f.write(json.dumps(c, ensure_ascii=False) + "\n")
                write_text(EXTRAKT / f"{bid}.vorschau.md", vorschau(b, rel, cells))
                eintrag["blaetter"].append(b)
                zaehler["extrahiert"] += 1

        for b in eintrag["blaetter"]:
            if b["meta"]["zellen"] == 0:
                b["duplikat_von"] = None
                b["leer"] = True
                continue
            if b["inhalt_hash"] in inhalt_index:
                b["duplikat_von"] = inhalt_index[b["inhalt_hash"]]
                zaehler["duplikat_blatt"] += 1
            else:
                inhalt_index[b["inhalt_hash"]] = b["blatt_id"]
        neu[fid] = eintrag

    for fid, e in alt.items():
        if fid not in neu:
            e["fehlt_im_input"] = True
            neu[fid] = e
            zaehler["fehlt"] += 1

    manifest["dateien"] = neu
    manifest["stand"] = now()
    save_yaml(MANIFEST, manifest)

    blaetter = sum(len(e.get("blaetter", [])) for e in neu.values())
    gruppen = Counter(b["profil"] for e in neu.values() for b in e.get("blaetter", [])
                      if not b.get("duplikat_von") and not b.get("leer"))
    eindeutig = sum(1 for e in neu.values() if not e.get("duplikat_von") and not e.get("fehlt_im_input"))
    print(f"Dateien: {len(dateien)}  eindeutig: {eindeutig}  Blätter: {blaetter}  "
          f"Layout-Gruppen: {len(gruppen)}")
    for k, v in zaehler.items():
        print(f"  {k}: {v}")
    append_log(f"[skript] p0_extract: {len(dateien)} Dateien, {blaetter} Blätter, "
               f"{len(gruppen)} Layout-Gruppen, {dict(zaehler)}")

    import queue_cli
    queue_cli.sync(verbose=True)


if __name__ == "__main__":
    main()
